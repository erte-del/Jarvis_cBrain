"""Slides: builds a deck from the teacher's template into a temp folder; nothing is opened.
    .venv/bin/python -m unittest tests.test_slides
"""

import asyncio
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from pptx import Presentation

from tools import registry, slides

DECK = {
    "title": "Logic gates", "unit": "Fundamentals of computer systems", "topic": "2",
    "slides": [
        {"kind": "objectives", "points": ["Knowledge:", "- Know the AND, OR and NOT gates"]},
        {"title": "AND gate", "points": ["Both inputs must be 1"], "table": [["A", "B", "Q"], ["0", "0", "0"], ["1", "1", "1"]],
         "notes": "What happens with 1 and 0?"},
        {"kind": "practice", "title": "Practice", "points": ["1. What is 1 AND 0?"], "notes": "1. 0"},
    ],
}


@unittest.skipUnless(slides.TEMPLATE.exists(), "no slide template in storage")
class SlidesTest(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        for patch in (mock.patch.object(slides.config, "FILES_DIR", Path(tmp.name)),
                      mock.patch.object(slides.subprocess, "run")):
            patch.start()
            self.addCleanup(patch.stop)
        self.out = Path(tmp.name) / "Output"

    def make(self):
        return asyncio.run(slides.make_slides.handler(DECK))["content"][0]["text"]

    def test_runs_freely(self):
        self.assertEqual(registry.classify("mcp__jarvis__make_slides"), "read")

    def test_builds_the_deck(self):
        self.assertIn("(5 slides)", self.make())
        deck = Presentation(str(self.out / "Logic gates.pptx"))
        names = [s.slide_layout.name for s in deck.slides]
        self.assertEqual(names, ["Title Slide", "7_Custom Layout", "3_Custom Layout", "2_Custom Layout", "8_Custom Layout"])
        self.assertIn("Logic gates", deck.slides[0].placeholders[11].text_frame.text)
        gate = deck.slides[2]
        self.assertEqual(gate.placeholders[13].text_frame.text, "AND gate")
        self.assertEqual([sh.table.cell(2, 2).text for sh in gate.shapes if sh.has_table], ["1"])
        self.assertEqual(gate.notes_slide.notes_text_frame.text, "What happens with 1 and 0?")
        footer = [sh.text_frame.text for sh in gate.slide_layout.shapes if sh.name.startswith("TextBox")]
        self.assertEqual(footer, ["Fundamentals of computer systems\nTopic 2 Logic gates"])

    def test_never_overwrites(self):
        self.make()
        self.make()
        self.assertTrue((self.out / "Logic gates 2.pptx").exists())

    def test_lectures_reads_a_deck(self):
        self.make()
        with mock.patch.object(slides, "LECTURES", self.out):
            listing = asyncio.run(slides.lectures.handler({}))["content"][0]["text"]
            text = asyncio.run(slides.lectures.handler({"deck": "logic"}))["content"][0]["text"]
        self.assertIn("Logic gates", listing)
        self.assertIn("--- Slide 3 (3_Custom Layout)", text)
        self.assertIn("A | B | Q", text)
        self.assertIn("Notes: What happens with 1 and 0?", text)
        self.assertEqual(registry.classify("mcp__jarvis__lectures"), "read")

    def test_lectures_pictures_skip_hidden_slides(self):
        self.make()
        path = self.out / "Logic gates.pptx"
        deck = Presentation(str(path))
        deck.slides[1]._element.set("show", "0")  # Keynote draws no picture for it
        deck.save(str(path))
        pictures = []
        for n in (1, 3, 4, 5):
            pictures.append(self.out / f"{n}.jpeg")
            pictures[-1].write_bytes(bytes([n]))
        with mock.patch.object(slides, "LECTURES", self.out), mock.patch.object(slides, "render", return_value=pictures):
            content = asyncio.run(slides.lectures.handler({"deck": "logic", "slides": [2, 3]}))["content"]
        self.assertEqual([c["type"] for c in content], ["text", "text", "text", "image"])
        self.assertIn("hidden", content[1]["text"])
        self.assertEqual(content[3]["data"], "Aw==")  # slide 3's picture


if __name__ == "__main__":
    unittest.main()
