"""Slides: PowerPoint lessons in the style of the user's computing teacher.

The template is the teacher's own deck (PG Online's AQA design) with its slides taken out,
kept in storage/slides/template.pptx (not in git: it's PG Online's). Each deck is built
from its layouts: title, hook question, objectives, teaching slides, activities, end.
Decks are saved in the Output folder of Jarvis's files and opened.
"""

import asyncio
import base64
import subprocess
from pathlib import Path
from typing import Any

from claude_agent_sdk import tool
from pptx import Presentation
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.util import Emu, Pt

import config

TEMPLATE = config.STORAGE_DIR / "slides" / "template.pptx"
LECTURES = Path.home() / "Desktop" / "school" / "computing" / "class lectures"
RENDERS = config.STORAGE_DIR / "lectures"
MAX_SEEN = 6  # slide pictures per call
COURSE = ["AQA", "AS Level", "Computer Science", "Paper 2"]
# slide kind: the teacher's layout for it
LAYOUTS = {"hook": "4_Custom Layout", "objectives": "7_Custom Layout", "content": "3_Custom Layout",
           "practice": "2_Custom Layout"}
LINE = Emu(420_000)  # height of one line of 24pt body text
ROW = Emu(380_000)
_keynote = asyncio.Lock()


def _fill(frame, points: list[str]) -> None:
    """Bullets; '- ' makes a sub-point and a line ending in ':' is a bold heading."""
    frame.text = ""
    for i, point in enumerate(points):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        sub = point.startswith("- ")
        para.level = 1 if sub else 0
        run = para.add_run()
        run.text = point[2:] if sub else point
        run.font.bold = point.rstrip().endswith(":") or None
        if sub:  # the layout greys sub-points out; the teacher makes them black
            run.font.color.theme_color = MSO_THEME_COLOR.TEXT_1
            para.space_after = Pt(12)


def _set(shape, lines: list[str]) -> None:
    """Replace a template text box's lines, keeping each line's formatting."""
    paras = shape.text_frame.paragraphs
    for para, line in zip(paras, lines):
        for run in para.runs[1:]:
            run._r.getparent().remove(run._r)
        if para.runs:
            para.runs[0].text = line
        else:
            para.add_run().text = line


def build(args: dict[str, Any], path: Path) -> int:
    deck = Presentation(str(TEMPLATE))
    title, unit, topic = args["title"], str(args.get("unit") or ""), str(args.get("topic") or "")
    # The layouts carry the topic: the number in the hexagon and the footer on every slide.
    for layout in deck.slide_layouts:
        for shape in layout.shapes:
            if shape.is_placeholder or not shape.has_text_frame or not shape.text_frame.text.strip():
                continue
            if shape.name.startswith("Hexagon"):
                _set(shape, [topic])
            else:
                _set(shape, [unit, f"Topic {topic} {title}" if topic else title])

    first = deck.slides.add_slide(deck.slide_layouts.get_by_name("Title Slide"))
    box = first.placeholders[10]
    box.left, box.top, box.width, box.height = box.left, box.top, Emu(2_750_617), box.height  # as wide as the teacher's
    box.text_frame.text = ""
    for i, line in enumerate(args.get("course") or COURSE):
        para = box.text_frame.paragraphs[0] if i == 0 else box.text_frame.add_paragraph()
        para.level = [0, 2, 3, 3][min(i, 3)]  # the layout's styles: AQA, AS Level, subject, paper
        run = para.add_run()
        run.text = line
        if i > 1:
            run.font.size = Pt(24)
    _fill(first.placeholders[11].text_frame, [title, "", unit])

    for s in args.get("slides") or []:
        kind = s.get("kind", "content")
        slide = deck.slides.add_slide(deck.slide_layouts.get_by_name(LAYOUTS.get(kind, LAYOUTS["content"])))
        points = [str(p) for p in s.get("points") or []]
        slide.placeholders[13].text = s.get("title") or ("Objectives" if kind == "objectives" else "")
        _fill(slide.placeholders[14].text_frame, points)
        rows = s.get("table") or []
        if rows:
            body = slide.placeholders[14]
            # A placeholder inherits its box; set all of it, or the width becomes 0.
            body.left, body.top, body.width, body.height = body.left, body.top, body.width, LINE * max(len(points), 1)
            cols = max(len(r) for r in rows)
            top = body.top + body.height + Emu(150_000)
            table = slide.shapes.add_table(len(rows), cols, body.left, top, body.width, ROW * len(rows)).table
            for r, row in enumerate(rows):
                for c in range(cols):
                    cell = table.cell(r, c)
                    cell.text = str(row[c]) if c < len(row) else ""
                    for para in cell.text_frame.paragraphs:
                        for run in para.runs:
                            run.font.size = Pt(18)
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]

    deck.slides.add_slide(deck.slide_layouts.get_by_name("8_Custom Layout"))
    deck.save(str(path))
    return len(deck.slides)


def _free(name: str) -> Path:
    out = config.FILES_DIR / "Output"
    out.mkdir(parents=True, exist_ok=True)
    stem = "".join(c for c in name if c not in '/\\:*?"<>|').strip() or "Slides"
    path, n = out / f"{stem}.pptx", 2
    while path.exists():
        path, n = out / f"{stem} {n}.pptx", n + 1
    return path


@tool(
    "make_slides",
    "Make a PowerPoint (.pptx) lesson in the user's computing teacher's design (AQA, PG Online "
    "template), save it in Jarvis's Output folder and open it. Give the whole deck at once. The "
    "title slide and the end slide are added for you. Each slide: kind (hook: an opening question "
    "in points; objectives: 'Knowledge:' and 'Skills:' headings with '- ' points under them; "
    "content: a teaching slide; practice: questions for the user to solve), title, points (short bullets; "
    "'- ' starts a sub-point, a line ending in ':' is a bold heading), optional table (rows of "
    "cells, first row the header), optional notes (speaker notes).",
    {
        "type": "object",
        "properties": {
            "title": {"type": "string", "description": "the lesson, e.g. 'Binary arithmetic'"},
            "unit": {"type": "string", "description": "the spec section, e.g. 'Data representation'"},
            "topic": {"type": "string", "description": "the topic number in the unit, e.g. '3'"},
            "course": {"type": "array", "items": {"type": "string"},
                       "description": f"title slide lines; default {COURSE}"},
            "slides": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "kind": {"type": "string", "enum": list(LAYOUTS)},
                        "title": {"type": "string"},
                        "points": {"type": "array", "items": {"type": "string"}},
                        "table": {"type": "array", "items": {"type": "array", "items": {"type": "string"}}},
                        "notes": {"type": "string"},
                    },
                },
            },
        },
        "required": ["title", "slides"],
    },
)
async def make_slides(args: dict[str, Any]) -> dict[str, Any]:
    if not TEMPLATE.exists():
        return {"content": [{"type": "text", "text": f"The slide template is missing: {TEMPLATE}"}], "is_error": True}
    path = _free(str(args.get("title") or ""))
    try:
        count = await asyncio.to_thread(build, args, path)
    except (KeyError, ValueError, TypeError) as e:
        return {"content": [{"type": "text", "text": f"Couldn't build the deck: {e!r}"}], "is_error": True}
    subprocess.run(["open", str(path)], check=False)
    return {"content": [{"type": "text", "text": f"Saved and opened {path} ({count} slides)."}]}


def read_deck(path: Path) -> list[str]:
    """Each slide's text: titles and bullets, tables as rows, then speaker notes."""
    out = []
    for n, slide in enumerate(Presentation(str(path)).slides, 1):
        hidden = " hidden" if slide._element.get("show") == "0" else ""
        lines = [f"--- Slide {n} ({slide.slide_layout.name}{hidden})"]
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                lines += ["  " * p.level + p.text for p in shape.text_frame.paragraphs if p.text.strip()]
            elif shape.has_table:
                lines += [" | ".join(c.text for c in row.cells) for row in shape.table.rows]
            elif shape.shape_type == 13:  # picture
                lines.append("[picture]")
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            lines.append("Notes: " + slide.notes_slide.notes_text_frame.text.strip())
        out.append("\n".join(lines))
    return out


EXPORT = """on run argv
  tell application "Keynote"
    repeat 60 times
      if exists document (item 1 of argv) then exit repeat
      delay 1
    end repeat
    set d to document (item 1 of argv)
    export d to POSIX file (item 2 of argv) as slide images with properties ¬
      {image format:JPEG, compression factor:0.8}
    close d saving no
  end tell
end run"""


def render(path: Path) -> list[Path]:
    """Every shown slide as a picture (Keynote leaves hidden ones out), drawn by Keynote once
    and kept in storage/lectures/.
    Keynote opens the deck by itself (AppleScript's own open can't read the Desktop)."""
    out = RENDERS / path.stem
    if not out.exists() or out.stat().st_mtime < path.stat().st_mtime:
        out.mkdir(parents=True, exist_ok=True)
        for old in out.glob("*.jpeg"):
            old.unlink()
        subprocess.run(["open", "-g", "-a", "Keynote", str(path)], check=True)
        subprocess.run(["osascript", "-", path.stem, str(out)], input=EXPORT.encode(),
                       capture_output=True, check=True, timeout=180)
        out.touch()
    return sorted(out.glob("*.jpeg"))


@tool(
    "lectures",
    "The user's computing teacher's own lesson decks (AQA Computer Science, PowerPoint). Without "
    "deck: the list. With deck (part of a file name, e.g. 'Topic 3' or 'sound'): the text of every "
    "slide, with its layout, tables and speaker notes. With slides too (slide numbers, up to "
    f"{MAX_SEEN}): those slides as pictures, to see their diagrams, images and layout; the first "
    "time for a deck takes a minute. Read one before make_slides as the example to follow.",
    {
        "type": "object",
        "properties": {
            "deck": {"type": "string"},
            "slides": {"type": "array", "items": {"type": "integer"}},
        },
    },
)
async def lectures(args: dict[str, Any]) -> dict[str, Any]:
    decks = sorted(LECTURES.glob("*.pptx"))
    want = str(args.get("deck") or "").strip().lower()
    hits = [d for d in decks if want in d.stem.lower()] if want else []
    if not hits:
        names = "\n".join(d.stem for d in decks) or f"none in {LECTURES}"
        return {"content": [{"type": "text", "text": f"The teacher's decks:\n{names}"}]}
    deck, seen = hits[0], [int(n) for n in args.get("slides") or []][:MAX_SEEN]
    try:
        texts = await asyncio.to_thread(read_deck, deck)
        if not seen:
            return {"content": [{"type": "text", "text": deck.stem + "\n" + "\n".join(texts)}]}
        seen = [n for n in seen if 1 <= n <= len(texts)]
        async with _keynote:
            pictures = await asyncio.to_thread(render, deck)
    except (OSError, ValueError, KeyError, subprocess.SubprocessError) as e:
        return {"content": [{"type": "text", "text": f"Couldn't read {deck.name}: {e!r}"}], "is_error": True}
    content: list[dict[str, Any]] = [{"type": "text", "text": deck.stem}]
    shown = [n for n, text in enumerate(texts, 1) if " hidden)" not in text.split("\n")[0]]
    for n in seen:
        content.append({"type": "text", "text": texts[n - 1]})
        if n in shown and shown.index(n) < len(pictures):
            data = base64.b64encode(pictures[shown.index(n)].read_bytes()).decode()
            content.append({"type": "image", "data": data, "mimeType": "image/jpeg"})
    return {"content": content}
